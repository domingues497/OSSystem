import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.commissions.models import DuplicataAssinatura


def _normalize_header(value):
    s = str(value or "").strip()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    s = s.upper().replace(" ", "").replace("-", "_")
    return s


def _to_int(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int,)):
        return int(value)
    if isinstance(value, (float,)):
        try:
            return int(value)
        except Exception:
            return None
    s = str(value).strip()
    if not s:
        return None
    s = s.replace(".", "").replace(",", "")
    if s.isdigit():
        try:
            return int(s)
        except Exception:
            return None
    try:
        return int(float(s))
    except Exception:
        return None


def _to_decimal(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return Decimal(int(value))
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except Exception:
            return None
    s = str(value).strip()
    if not s:
        return None
    s = s.replace("R$", "").replace(" ", "")
    if s.count(",") == 1 and s.count(".") >= 1:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _to_bool(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return bool(value)
    if isinstance(value, (int, float)):
        return bool(int(value))
    s = str(value).strip().lower()
    if not s:
        return None
    if s in {"1", "true", "t", "yes", "y", "sim", "s", "ok"}:
        return True
    if s in {"0", "false", "f", "no", "n", "nao", "não"}:
        return False
    return None


def _to_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


class Command(BaseCommand):
    help = "Importa planilha .xlsx para a tabela de controle de assinatura de duplicatas (commissions.DuplicataAssinatura)."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str)
        parser.add_argument("--sheet", type=str, default="")
        parser.add_argument("--truncate", action="store_true")
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument("--overwrite-assinada", action="store_true")

    def handle(self, *args, **options):
        xlsx_path = str(options["xlsx_path"])
        sheet_name = str(options.get("sheet") or "").strip()
        truncate = bool(options.get("truncate"))
        dry_run = bool(options.get("dry_run"))
        overwrite_assinada = bool(options.get("overwrite_assinada"))

        try:
            from openpyxl import load_workbook
        except Exception as e:
            raise CommandError("Dependência ausente: instale openpyxl (pip install openpyxl).") from e

        try:
            wb = load_workbook(filename=xlsx_path, data_only=True)
        except Exception as e:
            raise CommandError(f"Não foi possível abrir o arquivo: {xlsx_path}") from e

        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise CommandError(f"Aba '{sheet_name}' não encontrada. Abas: {', '.join(wb.sheetnames)}")
            ws = wb[sheet_name]
        else:
            ws = wb[wb.sheetnames[0]]

        header_row = None
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            if row and any(v is not None and str(v).strip() != "" for v in row):
                header_row = row
                break
        if not header_row:
            raise CommandError("Não foi possível identificar a linha de cabeçalho na planilha.")

        headers = [_normalize_header(h) for h in header_row]
        col = {h: i for i, h in enumerate(headers) if h}

        required = ["ESTAB", "DOCTO", "PARCELA", "SEQITEM"]
        missing = [h for h in required if h not in col]
        if missing:
            raise CommandError(f"Colunas obrigatórias ausentes: {', '.join(missing)}")

        def get(row_values, name):
            idx = col.get(name)
            if idx is None:
                return None
            if idx >= len(row_values):
                return None
            return row_values[idx]

        created = 0
        updated = 0
        skipped = 0

        with transaction.atomic():
            if truncate:
                if not dry_run:
                    DuplicataAssinatura.objects.all().delete()

            start_row_index = None
            for i in range(1, ws.max_row + 1):
                values = [c.value for c in ws[i]]
                if [_normalize_header(v) for v in values] == headers:
                    start_row_index = i + 1
                    break
            if not start_row_index:
                start_row_index = 2

            for row in ws.iter_rows(min_row=start_row_index, values_only=True):
                if not row or not any(v is not None and str(v).strip() != "" for v in row):
                    continue

                estab = _to_int(get(row, "ESTAB"))
                docto = _to_int(get(row, "DOCTO"))
                parcela = _to_int(get(row, "PARCELA"))
                seqitem = _to_int(get(row, "SEQITEM"))

                if estab is None or docto is None or parcela is None or seqitem is None:
                    skipped += 1
                    continue

                parsed_assinada = _to_bool(get(row, "ASSINADA"))

                defaults = {
                    "pessoa": str(get(row, "PESSOA") or "").strip(),
                    "cnpjf": _to_int(get(row, "CNPJF")),
                    "local": str(get(row, "LOCAL") or "").strip(),
                    "representante": str(get(row, "REPRESENTANTE") or "").strip(),
                    "dtemissao": _to_date(get(row, "DTEMISSAO")),
                    "dtvencto": _to_date(get(row, "DTVENCTO")),
                    "item": str(get(row, "ITEM") or "").strip(),
                    "grupocomissao": str(get(row, "GRUPOCOMISSAO") or "").strip(),
                    "comissaopercent": _to_decimal(get(row, "COMISSAOPERCENT")),
                    "tabprc": _to_int(get(row, "TABPRC")),
                    "tabela": str(get(row, "TABELA") or "").strip(),
                    "vlr_proporcional_item": _to_decimal(get(row, "VLR_PROPORCIONAL_ITEM")),
                    "previsao_comissao": _to_decimal(get(row, "PREVISAO_COMISSAO")),
                    "comissao_paga": _to_decimal(get(row, "COMISSAO_PAGA")),
                }

                if dry_run:
                    created += 1
                    continue

                obj, is_created = DuplicataAssinatura.objects.get_or_create(
                    estab=estab,
                    docto=docto,
                    parcela=parcela,
                    seqitem=seqitem,
                    defaults={**defaults, "assinada": bool(parsed_assinada) if parsed_assinada is not None else False},
                )
                if is_created:
                    created += 1
                    continue

                for k, v in defaults.items():
                    setattr(obj, k, v)
                if overwrite_assinada:
                    obj.assinada = bool(parsed_assinada) if parsed_assinada is not None else False
                else:
                    if parsed_assinada is not None:
                        obj.assinada = bool(parsed_assinada)
                obj.save()
                updated += 1

            if dry_run:
                transaction.set_rollback(True)

        self.stdout.write(
            self.style.SUCCESS(
                f"Import concluído. Criados: {created} | Atualizados: {updated} | Ignorados: {skipped} | dry_run={dry_run}"
            )
        )

